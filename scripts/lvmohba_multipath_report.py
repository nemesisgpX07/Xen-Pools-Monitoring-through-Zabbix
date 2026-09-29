#!/usr/bin/env python3
"""
File: lvmohba_multipath_report.py
Description:
    Scans all lvmohba SRs across one or more XenServer pools.
    Checks path counts from THREE locations (whichever the toolstack uses):
      1. PBD.other_config   – keys like  mpath-<SCSIid>
      2. host.other_config  – keys like  mpath-<SCSIid>   (XS 7.x / 8.x common)
      3. host.other_config  – key        multipath-<SCSIid>  (older format)
    Flags every (SR x Host) where active_paths < max_paths.

Usage:
    python3 lvmohba_multipath_report.py <host_id_1> [host_id_2 ...]
"""

import XenAPI
import sys, ssl, re, json, http.client, xmlrpc.client
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, GradientFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("FATAL: pip3 install openpyxl", file=sys.stderr); sys.exit(1)

try:
    from pass_config import get_xen_config
except ImportError:
    print("FATAL: pass_config.py not found.", file=sys.stderr); sys.exit(1)


# ── SSL (unchanged) ──────────────────────────────────────────────────────────

def setup_ssl_transport():
    ctx = ssl.create_default_context()
    ctx.check_hostname = False
    ctx.verify_mode = ssl.CERT_NONE
    class _Conn(http.client.HTTPSConnection):
        def __init__(self, h, **kw): super().__init__(h, context=ctx, **kw)
    def _make(self, host):
        if hasattr(self,'_connection') and self._connection and self._connection[0]==host:
            return self._connection[1]
        self._connection = host, _Conn(host); return self._connection[1]
    xmlrpc.client.SafeTransport.make_connection = _make


# ── Path-count parsers ───────────────────────────────────────────────────────

def _parse(raw):
    if not raw:
        return None, None
    s = str(raw).strip()

    if s.startswith(('{','[')):
        try:
            d = json.loads(s)
            if isinstance(d, list) and len(d) >= 2:
                return int(d[0]), int(d[1])
            if isinstance(d, dict):
                a = d.get('active_paths') or d.get('active') or d.get('active_by_host')
                t = d.get('max_paths')    or d.get('total')  or d.get('max')
                if isinstance(a, dict): a = sum(a.values())
                if a is not None and t is not None:
                    return int(a), int(t)
        except Exception:
            pass

    m = re.search(r'(\d+)\s*(?:/|of|:)\s*(\d+)', s)
    if m:
        return int(m.group(1)), int(m.group(2))

    return None, None


# ── Build a  { scsi_id: (active, max) }  map for one host ───────────────────

def _host_path_map(host_rec):
    result = {}
    oc = host_rec.get('other_config', {})
    for k, v in oc.items():
        kl = k.lower()
        for prefix in ('mpath-', 'multipath-'):
            if kl.startswith(prefix):
                scsi = k[len(prefix):]
                a, t = _parse(v)
                if t is not None:
                    if scsi not in result or (a is not None and a < (result[scsi][0] or 9999)):
                        result[scsi] = (a, t)
    return result


# ── Main detection ───────────────────────────────────────────────────────────

def find_multipath_failures(session):
    pools  = session.xenapi.pool.get_all_records()
    hosts  = session.xenapi.host.get_all_records()
    srs    = session.xenapi.SR.get_all_records()
    pbds   = session.xenapi.PBD.get_all_records()

    pool_rec         = list(pools.values())[0]
    pool_name        = pool_rec['name_label']
    master_host_name = hosts[pool_rec['master']]['name_label']

    lvmohba = {r: v for r, v in srs.items() if v.get('type','').lower() == 'lvmohba'}
    if not lvmohba:
        print(f"  [INFO] No lvmohba SRs in '{pool_name}'.", file=sys.stderr)
        return []

    host_path_maps = {ref: _host_path_map(rec) for ref, rec in hosts.items()}

    sr_pbds = {}
    for pbd_rec in pbds.values():
        sr_ref = pbd_rec.get('SR')
        if sr_ref in lvmohba:
            sr_pbds.setdefault(sr_ref, []).append(pbd_rec)

    failures = []

    for sr_ref, sr_rec in lvmohba.items():
        sr_name = sr_rec.get('name_label', '')
        sr_uuid = sr_rec.get('uuid', '')

        sr_scsi_ids = set()
        raw_scsi = sr_rec.get('sm_config', {}).get('devserial') or \
                   sr_rec.get('sm_config', {}).get('SCSIid')    or ''
        if raw_scsi:
            sr_scsi_ids.update(raw_scsi.replace('scsi-','').split(','))

        pbd_list = sr_pbds.get(sr_ref, [])

        for pbd_rec in pbd_list:
            host_ref  = pbd_rec.get('host')
            host_rec  = hosts.get(host_ref, {})
            host_name = host_rec.get('name_label', str(host_ref))
            attached  = pbd_rec.get('currently_attached', False)

            pbd_oc    = pbd_rec.get('other_config', {})
            dev_cfg   = pbd_rec.get('device_config', {})

            pbd_scsi = dev_cfg.get('SCSIid') or pbd_oc.get('SCSIid') or ''
            local_scsi_ids = set(sr_scsi_ids)
            if pbd_scsi:
                local_scsi_ids.add(pbd_scsi.replace('scsi-',''))

            pbd_mpath_keys = {k: v for k, v in pbd_oc.items()
                              if k.lower().startswith(('mpath-', 'multipath-'))}

            host_map = host_path_maps.get(host_ref, {})

            if local_scsi_ids:
                for hscsi in host_map:
                    for known in local_scsi_ids:
                        if hscsi.startswith(known[:20]) or known.startswith(hscsi[:20]):
                            local_scsi_ids.add(hscsi)

            reported_scsi = set()

            for k, v in pbd_mpath_keys.items():
                for prefix in ('mpath-', 'multipath-'):
                    if k.lower().startswith(prefix):
                        scsi = k[len(prefix):]
                a_val, t_val = _parse(v)
                if t_val is not None and a_val is not None and a_val < t_val:
                    failures.append(_row(pool_name, master_host_name, sr_name,
                                        sr_uuid, scsi, host_name, attached,
                                        a_val, t_val, 'PBD.other_config'))
                    reported_scsi.add(scsi)

            for scsi, (a_val, t_val) in host_map.items():
                if scsi in reported_scsi:
                    continue
                belongs = not local_scsi_ids
                for known in local_scsi_ids:
                    if scsi.startswith(known[:20]) or known.startswith(scsi[:20]) or scsi == known:
                        belongs = True; break
                if not belongs:
                    continue
                if t_val is not None and a_val is not None and a_val < t_val:
                    failures.append(_row(pool_name, master_host_name, sr_name,
                                        sr_uuid, scsi, host_name, attached,
                                        a_val, t_val, 'host.other_config'))
                    reported_scsi.add(scsi)

    total = len(failures)
    print(f"  [INFO] Pool '{pool_name}': {total} multipath failure(s) found.", file=sys.stderr)
    return failures


def _row(pool, master, sr_name, sr_uuid, scsi, host, attached, active, total, source):
    return {
        'Pool_Name':      pool,
        'Master_Host':    master,
        'SR_Name':        sr_name,
        'SR_UUID':        sr_uuid,
        'SCSI_ID':        scsi,
        'Host_Name':      host,
        'PBD_Attached':   'Yes' if attached else 'No',
        'Active_Paths':   active,
        'Max_Paths':      total,
        'Paths_Missing':  total - active,
        'Data_Source':    source,
        'Failure_Detail': f'{active} of {total} paths active',
    }


# ── Excel writer ─────────────────────────────────────────────────────────────

HEADERS = [
    'Pool_Name', 'Master_Host', 'SR_Name', 'SR_UUID', 'SCSI_ID',
    'Host_Name', 'PBD_Attached', 'Active_Paths', 'Max_Paths',
    'Paths_Missing', 'Data_Source', 'Failure_Detail',
]

# ── Corporate Dark palette (navy / red / amber) ───────────────────────────────
NAV_DARK   = '1B3A6B'   # header background
NAV_MID    = '14305A'   # header bottom border / subtitle text
NAV_LIGHT  = 'D6E4F7'   # subtitle bar background
DEAD_ROW   = 'FFECEC'   # row fill — 0 paths active
DEAD_BADGE = 'C62828'   # red accent colour
WARN_ROW   = 'FFF3CD'   # row fill — 2+ paths missing
WARN_BADGE = 'E65100'   # amber-orange accent
LOW_ROW    = 'FFFDE7'   # row fill — 1 path missing
LOW_BADGE  = 'F9A825'   # yellow accent
ALT_ROW    = 'F4F7FC'   # alternating faint blue-grey
BASE_ROW   = 'FFFFFF'
BORDER_CLR = 'C8D8EE'   # thin cell border (navy family)


def _border(bottom_heavy=False):
    bot = Side(style='medium' if bottom_heavy else 'thin',
               color=NAV_MID if bottom_heavy else BORDER_CLR)
    s   = Side(style='thin', color=BORDER_CLR)
    return Border(left=s, right=s, top=s, bottom=bot)


def _fill(hex_color):
    return PatternFill('solid', fgColor=hex_color)


def _style_header(cell):
    cell.font      = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
    cell.fill      = _fill(NAV_DARK)
    cell.border    = _border(bottom_heavy=True)
    cell.alignment = Alignment(horizontal='center', vertical='center')


def _style_data(cell, row_fill, mono=False, bold=False, color='1A1A1A'):
    cell.fill   = row_fill
    cell.border = _border()
    if mono:
        cell.font      = Font(name='Consolas', size=9, color=color)
        cell.alignment = Alignment(horizontal='left', vertical='center')
    else:
        cell.font      = Font(name='Calibri', bold=bold, color=color, size=10)
        cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)


def _autowidth(ws, hdr_row, cap=58):
    for col in ws.iter_cols():
        best = 10
        for cell in col:
            if cell.row < hdr_row:
                continue
            try:
                v = len(str(cell.value or ''))
                if v > best:
                    best = v
            except Exception:
                pass
        ws.column_dimensions[
            get_column_letter(col[0].column)].width = min(best + 4, cap)


def _title_block(ws, title, col_count, subtitle=''):
    ws.append([title] + [''] * (col_count - 1))
    r = ws.max_row
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=col_count)
    c = ws.cell(r, 1)
    c.font      = Font(name='Calibri', bold=True, color='FFFFFF', size=14)
    c.fill      = _fill(NAV_DARK)
    c.alignment = Alignment(horizontal='left', vertical='center')
    c.border    = _border()
    ws.row_dimensions[r].height = 36

    ws.append([subtitle] + [''] * (col_count - 1))
    r2 = ws.max_row
    ws.merge_cells(start_row=r2, start_column=1, end_row=r2, end_column=col_count)
    c2 = ws.cell(r2, 1)
    c2.font      = Font(name='Calibri', color=NAV_MID, size=10)
    c2.fill      = _fill(NAV_LIGHT)
    c2.alignment = Alignment(horizontal='left', vertical='center')
    c2.border    = _border()
    ws.row_dimensions[r2].height = 18


def write_excel(filename, data):
    from collections import defaultdict
    ts_label = datetime.now().strftime('%d %b %Y  %H:%M')
    wb = openpyxl.Workbook()

    # ── Sheet 1: All failures ─────────────────────────────────────────────────
    ws = wb.active
    ws.title = 'Multipath Failures'
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAV_DARK

    _title_block(ws, '  LVMOHBA Multipath Failures', len(HEADERS),
                 subtitle=f'  Generated: {ts_label}    |    {len(data)} failure(s) detected')

    ws.append(HEADERS)
    hdr_r = ws.max_row
    for cell in ws[hdr_r]:
        _style_header(cell)
    ws.row_dimensions[hdr_r].height = 28

    uuid_cols = {HEADERS.index('SR_UUID') + 1, HEADERS.index('SCSI_ID') + 1}
    num_cols  = {HEADERS.index('Active_Paths') + 1, HEADERS.index('Max_Paths') + 1,
                 HEADERS.index('Paths_Missing') + 1}
    det_col   = HEADERS.index('Failure_Detail') + 1

    for row in data:
        active  = row.get('Active_Paths', 1)
        missing = row.get('Paths_Missing', 0)

        if active == 0:
            rf, accent = DEAD_ROW, DEAD_BADGE
        elif missing >= 2:
            rf, accent = WARN_ROW, WARN_BADGE
        else:
            rf, accent = LOW_ROW, LOW_BADGE

        row_fill = _fill(rf)
        ws.append([row.get(h, '') for h in HEADERS])
        dr = ws.max_row
        ws.row_dimensions[dr].height = 20

        for ci, cell in enumerate(ws[dr], start=1):
            _style_data(cell, row_fill, mono=(ci in uuid_cols),
                        bold=(active == 0), color='1A1A1A')
            if ci in num_cols:
                cell.alignment = Alignment(horizontal='center', vertical='center')
                if ci == HEADERS.index('Active_Paths') + 1 and active == 0:
                    cell.font = Font(name='Calibri', bold=True, color=DEAD_BADGE, size=11)
            if ci == det_col:
                cell.value = f'{active} / {row["Max_Paths"]} paths active'
                cell.font  = Font(name='Calibri', bold=(active == 0),
                                  color=accent, size=10)
                cell.alignment = Alignment(horizontal='left', vertical='center')

    ws.freeze_panes = f'A{hdr_r + 1}'
    _autowidth(ws, hdr_r)
    for col_name, w in [('Active_Paths', 13), ('Max_Paths', 11),
                         ('Paths_Missing', 14), ('PBD_Attached', 13)]:
        ws.column_dimensions[
            get_column_letter(HEADERS.index(col_name) + 1)].width = w

    # ── Sheet 2: Summary by SR ────────────────────────────────────────────────
    ws2 = wb.create_sheet('Summary by SR')
    ws2.sheet_view.showGridLines = False
    ws2.sheet_properties.tabColor = '204F8C'
    sh2 = ['Pool_Name', 'SR_Name', 'SR_UUID',
           'Hosts_Affected', 'Max_Paths_Expected', 'Affected_Host_List']

    sr_map = defaultdict(lambda: {'hosts': set(), 'max_paths': 'N/A',
                                  'pool': '', 'sr_uuid': ''})
    for r in data:
        sr_map[r['SR_Name']]['hosts'].add(r['Host_Name'])
        sr_map[r['SR_Name']]['pool']      = r['Pool_Name']
        sr_map[r['SR_Name']]['sr_uuid']   = r['SR_UUID']
        sr_map[r['SR_Name']]['max_paths'] = r['Max_Paths']

    _title_block(ws2, '  Summary by Storage Repository', len(sh2),
                 subtitle=f'  {len(sr_map)} SR(s) with path failures    |    {ts_label}')
    ws2.append(sh2)
    hdr_r2 = ws2.max_row
    for cell in ws2[hdr_r2]:
        _style_header(cell)
    ws2.row_dimensions[hdr_r2].height = 28

    for i, (sr_name, info) in enumerate(sorted(sr_map.items())):
        n = len(info['hosts'])
        ws2.append([info['pool'], sr_name, info['sr_uuid'],
                    n, info['max_paths'], ', '.join(sorted(info['hosts']))])
        dr = ws2.max_row
        ws2.row_dimensions[dr].height = 20
        rf = _fill(ALT_ROW if i % 2 else BASE_ROW)
        for ci, cell in enumerate(ws2[dr], start=1):
            _style_data(cell, rf, mono=(ci == sh2.index('SR_UUID') + 1))
            if ci == sh2.index('Hosts_Affected') + 1:
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.font = Font(name='Calibri', bold=True,
                                 color=DEAD_BADGE if n >= 3 else WARN_BADGE, size=11)

    ws2.freeze_panes = f'A{hdr_r2 + 1}'
    _autowidth(ws2, hdr_r2)
    ws2.column_dimensions[
        get_column_letter(sh2.index('Hosts_Affected') + 1)].width = 16

    # ── Sheet 3: Summary by Host ──────────────────────────────────────────────
    ws3 = wb.create_sheet('Summary by Host')
    ws3.sheet_view.showGridLines = False
    ws3.sheet_properties.tabColor = '2A6099'
    sh3 = ['Pool_Name', 'Host_Name', 'SRs_Affected', 'SR_List']

    hmap = defaultdict(lambda: {'srs': set(), 'pool': ''})
    for r in data:
        hmap[r['Host_Name']]['srs'].add(r['SR_Name'])
        hmap[r['Host_Name']]['pool'] = r['Pool_Name']

    _title_block(ws3, '  Summary by Host', len(sh3),
                 subtitle=f'  {len(hmap)} host(s) reporting path failures    |    {ts_label}')
    ws3.append(sh3)
    hdr_r3 = ws3.max_row
    for cell in ws3[hdr_r3]:
        _style_header(cell)
    ws3.row_dimensions[hdr_r3].height = 28

    for i, (hname, info) in enumerate(sorted(hmap.items())):
        n = len(info['srs'])
        ws3.append([info['pool'], hname, n, ', '.join(sorted(info['srs']))])
        dr = ws3.max_row
        ws3.row_dimensions[dr].height = 20
        rf = _fill(ALT_ROW if i % 2 else BASE_ROW)
        for ci, cell in enumerate(ws3[dr], start=1):
            _style_data(cell, rf)
            if ci == sh3.index('SRs_Affected') + 1:
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.font = Font(name='Calibri', bold=True, color=WARN_BADGE, size=11)

    ws3.freeze_panes = f'A{hdr_r3 + 1}'
    _autowidth(ws3, hdr_r3)
    ws3.column_dimensions[
        get_column_letter(sh3.index('SRs_Affected') + 1)].width = 14

    wb.save(filename)
    print(f"\n✅  Report saved: {filename}", file=sys.stderr)


# ── Entry point ──────────────────────────────────────────────────────────────

def main():
    if len(sys.argv) < 2:
        print(f"Usage: {sys.argv[0]} <host_id_1> [host_id_2 ...]", file=sys.stderr)
        sys.exit(1)

    setup_ssl_transport()
    all_failures = []

    for host_id in sys.argv[1:]:
        cfg = get_xen_config(host_id)
        if not cfg:
            print(f"⚠  Skipping '{host_id}': not in pass_config.py", file=sys.stderr)
            continue
        print(f"\n--- Processing pool/host: {cfg['host']} ---", file=sys.stderr)
        session = None
        try:
            session = XenAPI.Session(f"https://{cfg['host']}")
            session.login_with_password(cfg['username'], cfg['password'])
            all_failures.extend(find_multipath_failures(session))
        except Exception as e:
            print(f"❌  {host_id}: {e}", file=sys.stderr)
        finally:
            if session:
                try: session.logout()
                except: pass

    if all_failures:
        all_failures.sort(key=lambda r: (r['Pool_Name'], r['SR_Name'], r['Host_Name']))
        ts  = datetime.now().strftime('%Y%m%d_%H%M%S')
        fn  = f"LVMOHBA_Multipath_Report_{ts}.xlsx"
        write_excel(fn, all_failures)
        print(f"\n🔴  Total failures: {len(all_failures)}", file=sys.stderr)
    else:
        print("\n⚠  No failures found via standard keys.", file=sys.stderr)
        print("   Run multipath_debug.py to inspect raw XenAPI data.", file=sys.stderr)

if __name__ == '__main__':
    main()
