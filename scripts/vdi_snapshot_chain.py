#!/usr/bin/env python3

# File: vdi_chain_excel_report.py
# Description: Generates a multi-tab Excel report of ALL VDI chains, last leaf nodes,
#              detects Problematic/Stuck chains, and maps Corrupted VHD logs to affected VMs.
# Usage: python3 vdi_chain_excel_report.py <host_id_1> <host_id_2> ... <host_id_N>

import XenAPI
import sys
import os
import ssl
import http.client
import xmlrpc.client
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import Font
except ImportError:
    print("FATAL: 'openpyxl' module is required to generate Excel files.", file=sys.stderr)
    print("Please install it by running: pip3 install openpyxl", file=sys.stderr)
    sys.exit(1)

try:
    from pass_config import get_xen_config
except ImportError:
    print("FATAL: pass_config.py not found.", file=sys.stderr)
    sys.exit(1)

def setup_ssl_transport():
    ssl_context = ssl.create_default_context()
    ssl_context.check_hostname = False
    ssl_context.verify_mode = ssl.CERT_NONE

    class CustomHTTPSConnection(http.client.HTTPSConnection):
        def __init__(self, host, **kwargs):
            super().__init__(host, context=ssl_context, **kwargs)

    def custom_https_connection(self, host):
        if hasattr(self, '_connection') and self._connection and self._connection[0] == host:
            return self._connection[1]
        self._connection = host, CustomHTTPSConnection(host)
        return self._connection[1]
    xmlrpc.client.SafeTransport.make_connection = custom_https_connection


def parse_corrupted_logs(filepath="all_corrupted_logs.txt"):
    """Reads the corrupted logs file and extracts Short ID, Dates, and Hit Count."""
    corrupted_logs = []
    if not os.path.exists(filepath):
        print(f"Warning: '{filepath}' not found. Tab 4 will be empty.", file=sys.stderr)
        return corrupted_logs

    try:
        with open(filepath, 'r') as f:
            lines = f.readlines()
            for line in lines[2:]:  # Skip the first 2 header lines
                if not line.strip(): continue
                
                # Split by whitespace. Multiple spaces are treated as one delimiter.
                parts = line.split()
                
                # A standard line should have at least 9 parts because dates have spaces
                # Example: ['mumctr4c-ctx229', '*13f7f0b8', 'Apr', '22', '15:56:44', 'Mar', '22', '23:21:12', '12']
                if len(parts) >= 9:
                    hostname = parts[0]
                    short_vhd = parts[1].replace('*', '')  # Remove asterisk
                    first_seen = f"{parts[2]} {parts[3]} {parts[4]}"
                    last_seen = f"{parts[5]} {parts[6]} {parts[7]}"
                    hit_count = parts[-1]  # Last element
                    
                    corrupted_logs.append({
                        "log_hostname": hostname,
                        "short_vhd": short_vhd,
                        "first_seen": first_seen,
                        "last_seen": last_seen,
                        "hit_count": hit_count
                    })
        print(f"Loaded {len(corrupted_logs)} corrupted VHD records from log file.", file=sys.stderr)
    except Exception as e:
        print(f"Error parsing {filepath}: {e}", file=sys.stderr)
    
    return corrupted_logs


def generate_reports(session, corrupted_logs):
    pools = session.xenapi.pool.get_all_records()
    hosts = session.xenapi.host.get_all_records()
    vms = session.xenapi.VM.get_all_records()
    vbds = session.xenapi.VBD.get_all_records()
    vdis = session.xenapi.VDI.get_all_records()
    srs = session.xenapi.SR.get_all_records()

    pool_rec = list(pools.values())[0]
    pool_name = pool_rec['name_label']
    master_host_name = hosts[pool_rec['master']]['name_label']

    vdi_to_parent = {}
    parent_to_children = {}
    vdi_to_vm_map = {}
    
    vdi_uuid_to_rec = {rec['uuid']: rec for rec in vdis.values()}

    # 1. Map VDIs to VMs
    for vbd in vbds.values():
        if vbd['type'] == 'Disk' and vbd['VDI'] != 'OpaqueRef:NULL':
            vdi_uuid = vdis.get(vbd['VDI'], {}).get('uuid')
            vm = vms.get(vbd['VM'])

            if vdi_uuid and vm:
                state = "Active VM"
                if vm['is_a_snapshot']: state = "VM Snapshot"
                elif vm['is_a_template']: state = "Template"
                elif vm['is_control_domain']: state = "Control Domain"

                vdi_to_vm_map[vdi_uuid] = {
                    "vm_name": vm['name_label'],
                    "state": state
                }

    # 2. Build Parent-Child relationships
    for vdi_ref, vdi_rec in vdis.items():
        vdi_uuid = vdi_rec['uuid']
        parent_uuid = vdi_rec.get('sm_config', {}).get('vhd-parent')

        if parent_uuid:
            vdi_to_parent[vdi_uuid] = parent_uuid
            if parent_uuid not in parent_to_children:
                parent_to_children[parent_uuid] = []
            parent_to_children[parent_uuid].append(vdi_uuid)

    total_data = []
    last_chain_data = []
    leaf_nodes = []

    # 3. Generate Tab 1 and Tab 2
    for vdi_ref, vdi_rec in vdis.items():
        vdi_uuid = vdi_rec['uuid']
        parent_uuid = vdi_to_parent.get(vdi_uuid, "None")
        is_child = vdi_uuid in vdi_to_parent
        is_parent = vdi_uuid in parent_to_children

        if is_child or is_parent:
            sr_name = srs.get(vdi_rec.get('SR'), {}).get('name_label', 'Unknown SR')
            mapping = vdi_to_vm_map.get(vdi_uuid, {"vm_name": "Unattached/None", "state": "Orphaned"})

            row = {
                "Poolname": pool_name,
                "Masterhost": master_host_name,
                "VMname": mapping['vm_name'],
                "State": mapping['state'],
                "Diskname": vdi_rec['name_label'],
                "SRname": sr_name,
                "VDI_UUID": vdi_uuid,
                "Parent_UUID": parent_uuid
            }
            total_data.append(row)

            if not is_parent:
                last_chain_data.append(row)
                leaf_nodes.append(vdi_uuid)

    problematic_data = []
    corrupted_chain_data = []

    # 4. Generate Tab 3 & Tab 4
    for leaf in leaf_nodes:
        current = leaf
        depth = 1
        path_uuids = [leaf]

        while current in vdi_to_parent and vdi_to_parent[current] != "None":
            current = vdi_to_parent[current]
            path_uuids.append(current)
            depth += 1
            if depth > 50: break 

        root_uuid = path_uuids[-1]
        mapping = vdi_to_vm_map.get(leaf, {"vm_name": "Unattached/None", "state": "Orphaned"})
        leaf_vdi = vdi_uuid_to_rec.get(leaf, {})
        sr_name = srs.get(leaf_vdi.get('SR'), {}).get('name_label', 'Unknown SR')

        # --- Tab 3 (Mathematical Analysis) ---
        issues = []
        if depth >= 3: issues.append(f"Deep Chain (Depth {depth})")
        if mapping['state'] == "Orphaned": issues.append("Orphaned Chain")

        if issues:
            problematic_data.append({
                "Poolname": pool_name, "Masterhost": master_host_name,
                "End_Leaf_VM": mapping['vm_name'], "End_Leaf_State": mapping['state'],
                "End_Leaf_Disk": f"{leaf_vdi.get('name_label', 'Unknown')} ({leaf})",
                "SRname": sr_name, "Root_Parent_UUID": root_uuid,
                "Chain_Depth": depth, "Issue_Type": " + ".join(issues)
            })

        # --- Tab 4 (Corrupted Log Matches) ---
        for corrupt_log in corrupted_logs:
            short_id = corrupt_log['short_vhd']
            for chain_uuid in path_uuids:
                if chain_uuid.startswith(short_id):
                    # Found a matched VHD inside this chain!
                    corrupted_chain_data.append({
                        "Poolname": pool_name,
                        "Log_Hostname": corrupt_log['log_hostname'],
                        "Affected_Leaf_VM": mapping['vm_name'],
                        "Leaf_State": mapping['state'],
                        "Leaf_Disk_Name": leaf_vdi.get('name_label', 'Unknown'),
                        "Leaf_VDI_UUID": leaf,
                        "Leaf_Parent_UUID": vdi_to_parent.get(leaf, "None"),
                        "Corrupt_Short_VHD": short_id,
                        "Matched_Full_UUID": chain_uuid,
                        "Corrupted_Disk_Parent_UUID": vdi_to_parent.get(chain_uuid, "None"),
                        "First_Seen": corrupt_log['first_seen'],
                        "Last_Seen": corrupt_log['last_seen'],
                        "Hit_Count": corrupt_log['hit_count'],
                        "SRname": sr_name,
                        "Chain_Depth": depth
                    })
                    break 

    return total_data, last_chain_data, problematic_data, corrupted_chain_data

def write_excel(filename, total_data, last_chain_data, problematic_data, corrupted_chain_data):
    wb = openpyxl.Workbook()

    def create_sheet(sheet, title, data, headers):
        sheet.title = title
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)

        for row in data:
            sheet.append([row.get(h, "") for h in headers])

        for col in sheet.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except: pass
            sheet.column_dimensions[column].width = (max_length + 2)

    # Sheet 1
    ws1 = wb.active
    create_sheet(ws1, "Total Data", total_data, ["Poolname", "Masterhost", "VMname", "State", "Diskname", "SRname", "VDI_UUID", "Parent_UUID"])

    # Sheet 2
    ws2 = wb.create_sheet()
    create_sheet(ws2, "Last Chain Data", last_chain_data, ["Poolname", "Masterhost", "VMname", "State", "Diskname", "SRname", "VDI_UUID", "Parent_UUID"])

    # Sheet 3
    ws3 = wb.create_sheet()
    create_sheet(ws3, "Problematic Disks Trace", problematic_data, ["Poolname", "Masterhost", "End_Leaf_VM", "End_Leaf_State", "End_Leaf_Disk", "SRname", "Root_Parent_UUID", "Chain_Depth", "Issue_Type"])

    # Sheet 4 (UPDATED)
    ws4 = wb.create_sheet()
    headers_tab4 = [
        "Poolname", "Log_Hostname", "Affected_Leaf_VM", "Leaf_State", 
        "Leaf_Disk_Name", "Leaf_VDI_UUID", "Leaf_Parent_UUID", 
        "Corrupt_Short_VHD", "Matched_Full_UUID", "Corrupted_Disk_Parent_UUID", 
        "First_Seen", "Last_Seen", "Hit_Count", "SRname", "Chain_Depth"
    ]
    create_sheet(ws4, "Corrupted Chains", corrupted_chain_data, headers_tab4)

    wb.save(filename)
    print(f"\nExcel Report successfully saved to: {filename}", file=sys.stderr)


def main():
    if len(sys.argv) < 2:
        print(f"Usage: {sys.argv[0]} <host_id_1> <host_id_2> ... <host_id_N>", file=sys.stderr)
        sys.exit(1)

    host_ids = sys.argv[1:]
    corrupted_logs = parse_corrupted_logs("all_corrupted_logs.txt")

    master_total = []
    master_last_chain = []
    master_problematic = []
    master_corrupted = []

    setup_ssl_transport()

    for host_id in host_ids:
        config = get_xen_config(host_id)
        if not config:
            print(f"Skipping '{host_id}': No config found in pass_config.py", file=sys.stderr)
            continue

        print(f"\n--- Processing Host/Pool: {config['host']} ---", file=sys.stderr)
        session = None
        try:
            session = XenAPI.Session(f"https://{config['host']}")
            session.login_with_password(config['username'], config['password'])

            tot, last, prob, corr = generate_reports(session, corrupted_logs)
            
            master_total.extend(tot)
            master_last_chain.extend(last)
            master_problematic.extend(prob)
            master_corrupted.extend(corr)

        except Exception as exc:
            print(f"Error processing {host_id}: {exc}", file=sys.stderr)
        finally:
            if session and session.is_logged_in():
                session.logout()

    if master_total:
        filename = f"Master_VDI_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        write_excel(filename, master_total, master_last_chain, master_problematic, master_corrupted)
    else:
        print("\nNo data was collected. Excel file was not generated.")

if __name__ == "__main__":
    main()
